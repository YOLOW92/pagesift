#!/usr/bin/env python3
"""候选真值：用 pandoc 3.9（独立实现）把文件转成 HTML，取标题、表格、列表、图片替代文字、脚注、正文样句。
pptx 的演讲者备注和图片说明另用 python-pptx 读。输出只是候选，之后逐份人工核对、修改（见 truth/*/REVIEW.md）。
用法：~/.venvs/label/bin/python label.py <dev|holdout>"""
import html.parser, json, os, re, subprocess, sys, unicodedata
SET = sys.argv[1]
PANDOC = os.path.expanduser("~/.venvs/label/lib/python3.12/site-packages/pypandoc/files/pandoc")
CAP = {"headings": 30, "tables": 10, "lists": 20, "images": 10, "footnotes": 10, "notes": 10}
def norm(s): return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", s)).strip()

class H(html.parser.HTMLParser):
    def __init__(s):
        super().__init__(); s.headings=[]; s.tables=[]; s.lists=[]; s.images=[]; s.paras=[]; s.footnotes=[]
        s.stack=[]; s.buf=None; s.tab=[]; s.in_fn=0; s.notes=[]; s.in_notes=0
    def handle_starttag(s, t, a):
        a=dict(a)
        if t=="section" and "footnotes" in (a.get("class") or ""): s.in_fn+=1
        if t=="div" and "notes" in (a.get("class") or ""): s.in_notes+=1
        if t in ("h1","h2","h3","h4","h5","h6","li","p","td","th"):
            s.stack.append((t,a,[]))
        if t=="table": s.tab.append([])
        if t=="tr" and s.tab: s.tab[-1].append([])
        if t=="img" and norm(a.get("alt") or ""): s.images.append(norm(a["alt"]))
        if t=="br": s.handle_data(" ")
    def handle_endtag(s, t):
        if t=="section" and s.in_fn: s.in_fn-=1
        if t=="div" and s.in_notes: s.in_notes-=1
        if s.stack and s.stack[-1][0]==t:
            tag,a,buf=s.stack.pop(); text=norm("".join(buf))
            if tag[0]=="h" and tag[1:].isdigit() and text: s.headings.append([int(tag[1]), text])
            elif tag=="li" and text:
                (s.footnotes if s.in_fn else s.lists).append(text.removesuffix("↩︎").strip())
            elif tag=="p" and text:
                if not any(x[0] in ("li","td","th") for x in s.stack):
                    (s.notes if s.in_notes else s.paras).append(text)
            elif tag in ("td","th") and s.tab and s.tab[-1]:
                span=int(a.get("colspan") or 1); s.tab[-1][-1].extend([text]+[""]*(span-1))
        if t=="table" and s.tab:
            rows=s.tab.pop(); rows=[r for r in rows if r]
            if rows: s.tables.append(rows)
    def handle_data(s, d):
        for x in s.stack: x[2].append(d)

def samples(rows):
    cells=[c for r in rows for c in r if 0<len(c)<=60]
    pick=[]
    for c in cells[:2]+cells[len(cells)//2:len(cells)//2+1]+cells[-1:]:
        if c not in pick: pick.append(c)
    return pick

def spread(xs, n):
    if len(xs)<=n: return xs
    return [xs[i*len(xs)//n] for i in range(n)]

out_dir=f"truth/{SET}"; os.makedirs(out_dir, exist_ok=True)
for fmt in ("docx","xlsx","pptx","epub","csv"):
    d=f"{SET}/{fmt}"
    if not os.path.isdir(d): continue
    for name in sorted(os.listdir(d)):
        if not name.endswith("."+fmt): continue
        path=os.path.join(d,name)
        r=subprocess.run([PANDOC,"-f",fmt,"-t","html",path],capture_output=True,timeout=300)
        h=H(); h.feed(r.stdout.decode("utf-8","replace"))
        t={"file":name,"format":fmt,
           "headings":h.headings[:CAP["headings"]],
           "tables":[{"rows":len(rows),"cols":max(map(len,rows)),"cells":samples(rows)} for rows in h.tables[:CAP["tables"]]],
           "lists":[x[:80] for x in h.lists[:CAP["lists"]]],
           "images":h.images[:CAP["images"]],
           "footnotes":[x[:80] for x in h.footnotes[:CAP["footnotes"]]],
           "notes":[x[:80] for x in h.notes[:CAP["notes"]]],
           "text":[p[:80] for p in spread([p for p in h.paras if len(p)>=40],5)],
           "order":[x[1] for x in spread(h.headings,8)]}
        if fmt=="pptx":
            from pptx import Presentation
            pr=Presentation(path); notes=[]; alts=[]
            for sl in pr.slides:
                if sl.has_notes_slide:
                    nt=norm(sl.notes_slide.notes_text_frame.text if sl.notes_slide.notes_text_frame else "")
                    if nt: notes.append(nt[:80])
                for sh in sl.shapes:
                    try: descr=norm(sh._element._nvXxPr.cNvPr.get("descr") or "")
                    except Exception: descr=""
                    if descr and sh.shape_type==13: alts.append(descr[:80])
            t["notes"]=notes[:CAP["notes"]]; t["images"]=alts[:CAP["images"]]
        if r.returncode: t["pandoc_error"]=r.stderr.decode()[:300]
        json.dump(t,open(f"{out_dir}/{name}.json","w"),ensure_ascii=False,indent=1)
        print(name, {k:len(v) for k,v in t.items() if isinstance(v,list)}, t.get("pandoc_error",""))
